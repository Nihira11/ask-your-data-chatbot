"""Extract a reproducible small sample from UCI's original workbook (requires openpyxl).

Usage: python scripts/prepare_sample.py /path/to/online+retail.zip
Reads the workbook only. Does not create or edit spreadsheets.
"""
import csv
import hashlib
import json
import sys
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


def main():
    source = Path(sys.argv[1])
    destination = Path(__file__).resolve().parents[1] / 'data' / 'sample'
    with zipfile.ZipFile(source) as archive:
        with archive.open('Online Retail.xlsx') as stream:
            workbook = load_workbook(stream, read_only=True, data_only=True)
            sheet = workbook.active
            selected = defaultdict(set)
            rows = []
            for row in sheet.iter_rows(min_row=2, values_only=True):
                invoice, stock, description, quantity, date, price, customer, country = row
                if not invoice or not stock or date is None:
                    continue
                invoice = str(invoice)
                key = (date.strftime('%Y-%m'), invoice.upper().startswith('C'))
                limit = 5 if key[1] else 25
                if invoice not in selected[key] and len(selected[key]) >= limit:
                    continue
                selected[key].add(invoice)
                rows.append([invoice, str(stock), description or '', int(quantity), date.isoformat(),
                             price, str(int(customer)) if customer is not None else '', country or ''])
            workbook.close()
    target = destination / 'online_retail_sample.csv'
    with target.open('w', newline='', encoding='utf-8') as file:
        writer = csv.writer(file)
        writer.writerow(['invoice_id', 'product_id', 'description', 'quantity', 'invoice_date', 'unit_price', 'customer_id', 'country'])
        writer.writerows(rows)
    manifest = {'version': 'uci-retail-sample-v1', 'source': 'https://archive.ics.uci.edu/dataset/352/online-retail',
                'citation': 'Chen, D. (2015). Online Retail. UCI Machine Learning Repository. https://doi.org/10.24432/C5BW33',
                'license': 'CC BY 4.0', 'selection': 'First 25 non-cancellation and first 5 cancellation invoice IDs per calendar month in workbook order; all their lines retained.',
                'rows': len(rows), 'invoices': len({r[0] for r in rows}), 'date_min': min(r[4] for r in rows),
                'date_max': max(r[4] for r in rows), 'sha256': hashlib.sha256(target.read_bytes()).hexdigest(),
                'source_zip_sha256': hashlib.sha256(source.read_bytes()).hexdigest()}
    (destination / 'manifest.json').write_text(json.dumps(manifest, indent=2) + '\n')
    print(json.dumps(manifest, indent=2))


if __name__ == '__main__':
    main()
